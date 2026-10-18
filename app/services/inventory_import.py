"""Reading an Excel sheet of products. Plain Python: give it the file's bytes, get rows back.

The sheet's first row holds the column names, the same ones as the Add product form. Names are
matched loosely (capitals, spaces and punctuation don't matter), and a few common alternatives
work too, like "Qty" for Stock. Every sheet row becomes a dict of text values keyed by the
form's field names, ready for the same checks as the Add product form.
"""
import re
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_ROWS = 5000
MAX_BYTES = 2 * 1024 * 1024

# form field -> the heading shown in the template, then other names that mean the same thing
FIELDS = {
    "code": ("Item Code/Bar Code", "itemcode", "barcode", "code"),
    "name": ("Name", "productname", "itemname", "product", "item"),
    "category": ("Category",),
    "price": ("Selling Price", "price", "sellingpricerupees"),
    "stock": ("Stock", "quantity", "qty"),
    "low_stock_level": ("Low-stock level", "lowstock", "lowlevel", "reorderlevel"),
    "hsn": ("HSN/SAC number", "hsnsac", "hsn", "sac"),
    "supplier": ("Supplier", "suppliername", "vendor", "vendorname"),
    "cost_price": ("Cost price", "cost", "costpricerupees"),
    "gst_rate": ("GST rate", "gst", "gstpercent"),
}


class SheetError(Exception):
    """The file can't be read as a product sheet; the message is for the owner."""


def squash(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def lookup():
    names = {}
    for field, (heading, *others) in FIELDS.items():
        for n in (heading, *others):
            names[squash(n)] = field
    return names


def as_text(value):
    """A cell as the text someone would have typed: 36.0 -> '36', 299.5 -> '299.5'."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        d = Decimal(repr(value))
        return str(int(d)) if d == d.to_integral_value() else format(d, "f")
    return str(value).strip()


def read_products(data):
    """Returns a list of (sheet row number, {field: text}). Raises SheetError."""
    if len(data) > MAX_BYTES:
        raise SheetError("That file is too big. Keep it under 2 MB.")
    if data[:2] != b"PK":                       # an .xlsx file is a zip file underneath
        raise SheetError("That doesn't look like an Excel (.xlsx) file.")
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
    except Exception:
        raise SheetError("That Excel file couldn't be read. Is it an .xlsx file?")
    if not header:
        raise SheetError("The sheet is empty. The first row needs the column names.")
    names = lookup()
    columns = {}
    for index, cell in enumerate(header):
        field = names.get(squash(cell))
        if field and field not in columns.values():
            columns[index] = field
    missing = [FIELDS[f][0] for f in ("name", "price") if f not in columns.values()]
    if missing:
        raise SheetError("The first row needs a column called " + " and ".join(missing)
                         + ". Download the template to see the column names.")
    found = []
    for number, row in enumerate(rows, start=2):
        values = {field: as_text(row[i]) if i < len(row) else "" for i, field in columns.items()}
        if not any(values.values()):
            continue                            # a blank row
        found.append((number, values))
        if len(found) > MAX_ROWS:
            raise SheetError(f"Too many rows. Import up to {MAX_ROWS} products at a time.")
    if not found:
        raise SheetError("No products found under the column names.")
    return found


def template_bytes():
    """An empty sheet with the right column names and one example row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append([heading for heading, *_ in FIELDS.values()])
    ws.append(["KA36", "Cotton bedsheet", "Bedsheets", 1299, 20, 5, "6302", "Sai Textiles", 800, 5])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i in range(1, len(FIELDS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
