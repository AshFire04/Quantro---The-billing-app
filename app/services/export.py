"""The Excel download of bills. Plain Python: give it a database and bills, get the file's bytes."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .billing import PAYMENT_MODES

MONEY = '#,##0.00'


def bills_workbook(db, bills):
    """Two sheets: one row per bill, and one row per item sold. Amounts are in rupees."""
    ids = [b["id"] for b in bills]
    marks = ",".join("?" * len(ids))
    paid, lines = {}, {}
    if ids:
        for p in db.execute(f"SELECT bill_id, mode, amount FROM payments WHERE bill_id IN ({marks})", ids):
            mode = p["mode"] if p["mode"] in PAYMENT_MODES else "Other"
            paid.setdefault(p["bill_id"], {})
            paid[p["bill_id"]][mode] = paid[p["bill_id"]].get(mode, 0) + p["amount"]
        for l in db.execute(f"SELECT * FROM bill_lines WHERE bill_id IN ({marks}) ORDER BY id", ids):
            lines.setdefault(l["bill_id"], []).append(l)
    labels = {b["id"]: b["bill_label"] for b in bills}

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Bills"
    sheet.append(["Bill no", "Date", "Customer", "Phone", "Location", "State", "Customer type",
                  "Order came through", "Status", "Items total", "Whole-bill discount %",
                  "Whole-bill discount", "Shipping & handling", "Round off", "Total",
                  "UPI", "Cash", "Net-banking", "Other payment", "Due", "Offers applied",
                  "Replaces", "Replaced by"])
    money_columns = {10, 12, 13, 14, 15, 16, 17, 18, 19, 20}
    for b in bills:
        items_total = sum(l["line_amount"] for l in lines.get(b["id"], []))
        whole_off = items_total - (b["total"] - b["shipping"] - b["round_off"])
        by_mode = paid.get(b["id"], {})
        sheet.append([
            b["bill_label"], b["created_at"][:10], b["customer_name"], b["customer_phone"],
            b["customer_location"], b["place_of_sale"], b["customer_type"], b["order_source"],
            b["status"], items_total / 100, b["whole_discount_pct"], whole_off / 100,
            b["shipping"] / 100, b["round_off"] / 100, b["total"] / 100,
            by_mode.get("UPI", 0) / 100, by_mode.get("Cash", 0) / 100,
            by_mode.get("Net-banking", 0) / 100, by_mode.get("Other", 0) / 100,
            (b["due_amount"] if b["status"] == "Saved" else 0) / 100, b["offers_applied"],
            labels.get(b["replaces"], ""), labels.get(b["replaced_by"], "")])

    items = wb.create_sheet("Items")
    items.append(["Bill no", "Date", "Customer", "Item code", "Item", "Quantity", "Rate",
                  "Discount %", "Amount", "Status"])
    for b in bills:
        for l in lines.get(b["id"], []):
            items.append([b["bill_label"], b["created_at"][:10], b["customer_name"], l["code"] or "",
                          l["name"], l["quantity"], l["rate"] / 100, l["line_discount_pct"],
                          l["line_amount"] / 100, b["status"]])

    for ws, money in ((sheet, money_columns), (items, {7, 9})):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"        # always text, so '=...' typed by someone is never a formula
                elif cell.column in money:
                    cell.number_format = MONEY
        for number, column in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[get_column_letter(number)].width = min(max(width + 2, 10), 40)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def profit_workbook(figures, start="", end=""):
    """One sheet: a row per product, then a totals row. Amounts are in rupees."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Profit"
    ws.append(["Period", f"{start or 'start'} to {end or 'today'}"])
    ws.append([])
    ws.append(["Product", "Item code", "Supplier", "Cost price", "Selling price", "Profit per item",
               "Quantity sold", "Sales", "Cost", "Profit", "Margin %"])
    header_row = 3
    for r in figures["rows"]:
        ws.append([r["name"] + (" (removed)" if r["removed"] else ""), r["code"], r["supplier"],
                   r["cost_price"] / 100 if r["cost_price"] else None,
                   r["price"] / 100 if r["price"] else None,
                   r["per_item"] / 100 if r["per_item"] is not None else None,
                   r["qty"] or None, r["revenue"] / 100 if r["revenue"] else None,
                   r["cost"] / 100 if r["cost"] else None,
                   r["profit"] / 100 if r["revenue"] else None,
                   round(r["margin"], 1) if r["margin"] is not None else None])
    ws.append(["Total", "", "", None, None, None, sum(r["qty"] for r in figures["rows"]) or None,
               figures["revenue"] / 100, figures["cost"] / 100, figures["profit"] / 100,
               round(figures["margin"], 1) if figures["margin"] is not None else None])
    if figures["unknown_revenue"]:
        ws.append([])
        ws.append([f"Sales of {figures['unknown_revenue'] / 100:,.2f} are not counted: "
                   "those items have no cost price."])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for cell in ws[header_row + len(figures["rows"]) + 1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"            # text stays text, never a formula
            elif cell.column in (4, 5, 6, 8, 9, 10):
                cell.number_format = MONEY
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    for number, column in enumerate(ws.iter_cols(min_row=header_row, max_row=ws.max_row - 0), 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(number)].width = min(max(width + 2, 10), 40)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
