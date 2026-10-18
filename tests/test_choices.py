import sqlite3
from io import BytesIO

from openpyxl import load_workbook

from app.db import init_db
from app.services import choices
from tests.test_auth import create
from tests.test_inventory import add


def setup(client):
    create(client)
    add(client, name="Tray", price="450", stock="50")


def page(client, url):
    import html
    return html.unescape(client.get(url).data.decode())


def bill(client, source="Walk-in", mode="Cash", **extra):
    return client.post("/billing/save", json={
        "items": [{"product_id": 1, "quantity": 1}], "customer": "Asha",
        "order_source": source, "payments": [{"mode": mode}], **extra})


def test_a_new_shop_starts_with_the_common_choices(client, app):
    setup(client)
    with app.app_context():
        db = sqlite3.connect(app.config["DATABASE"])
        db.row_factory = sqlite3.Row
        assert choices.names(db, 1, "source") == ["Walk-in", "WhatsApp", "Instagram", "Phone call"]
        assert choices.names(db, 1, "method") == ["Cash", "UPI", "Card", "Bank transfer"]


def test_shop_details_shows_both_lists_and_the_sections(client):
    setup(client)
    p = page(client, "/settings")
    for text in ("Shop profile", "Monthly target", "QR codes", "Offers", "Order source",
                 "Payment methods", "Order came through", "Payment method", "Bill preview",
                 "WhatsApp", "Bank transfer", "Add a source", "Add a method"):
        assert text in p


def test_adding_a_choice_shows_it_on_billing_and_past_bills(client):
    setup(client)
    client.post("/settings/sources", data={"name": "  Website "})
    client.post("/settings/payments", data={"name": "Credit"})
    billing = page(client, "/billing")
    assert "<option>Website</option>" in billing and "<option>Credit</option>" in billing
    assert "<option>Other</option>" in billing                      # always there as well
    assert bill(client, source="Website", mode="Credit").status_code == 200
    past = page(client, "/bills")
    assert "<th class=\"num\">Credit</th>" in past
    assert "<option >Website</option>" in past or "Website</option>" in past


def test_a_bill_can_only_use_choices_in_the_list_or_other(client):
    setup(client)
    res = bill(client, source="Carrier pigeon")
    assert res.status_code == 400 and "order source" in res.get_json()["error"]
    res = bill(client, mode="Barter")
    assert res.status_code == 400 and "payment mode" in res.get_json()["error"]
    ok = client.post("/billing/save", json={
        "items": [{"product_id": 1, "quantity": 1}], "customer": "Asha", "order_source": "Other",
        "order_source_other": "Pigeon", "payments": [{"mode": "Other", "other": "Barter"}]})
    assert ok.status_code == 200
    shown = page(client, "/bills/1")
    assert "Other: Pigeon" in shown and "Barter" in shown


def test_a_removed_choice_cannot_be_used_but_old_bills_keep_it(client):
    setup(client)
    assert bill(client, source="WhatsApp", mode="UPI").status_code == 200
    sources = chip_ids(page(client, "/settings"), "sources")
    client.post(f"/settings/sources/{sources['WhatsApp']}/delete")
    methods = chip_ids(page(client, "/settings"), "payments")
    client.post(f"/settings/payments/{methods['UPI']}/delete")
    assert bill(client, source="WhatsApp").status_code == 400
    assert bill(client, mode="UPI").status_code == 400
    old = page(client, "/bills/1")
    assert "WhatsApp" in old and "UPI" in old                       # the old bill is unchanged
    # still findable, and its payment still adds up in an Other column
    past = page(client, "/bills")
    assert "<th class=\"num\">Other</th>" in past
    assert "Asha" in page(client, "/bills?source=WhatsApp")
    assert "Asha" in page(client, "/bills?mode=UPI")
    # editing it brings the removed names back as 'Other' with the name typed in
    edit = page(client, "/billing/edit/1")
    assert '"source": "Other"' in edit and '"source_other": "WhatsApp"' in edit


def chip_ids(text, word):
    """{name: id} for the chips in a settings list, read from their remove-form addresses."""
    import re
    return {m.group(1): int(m.group(2)) for m in re.finditer(
        rf'<span>([^<]+)</span>\s*<form method="post" action="/settings/{word}/(\d+)/delete"', text)}


def test_choice_rules(client):
    setup(client)
    for word, name, message in [
            ("sources", "walk-in", "already in the list"), ("sources", " ", "type the order source"),
            ("sources", "x" * 41, "under 40"), ("sources", "Other", "always there"),
            ("payments", "cash", "already in the list"), ("payments", "other: cheque", "always there")]:
        res = client.post(f"/settings/{word}", data={"name": name}, follow_redirects=True)
        assert message in page_from(res)
    assert page(client, "/settings").count("WhatsApp") >= 1


def page_from(res):
    import html
    return html.unescape(res.data.decode())


def test_choices_are_private_to_each_shop_and_need_sign_in(client):
    assert client.post("/settings/sources", data={"name": "X"}).status_code == 302
    setup(client)
    client.post("/settings/sources", data={"name": "Mine"})
    client.get("/logout")
    create(client, login="two@example.com", shop="Other Shop")
    p = page(client, "/settings")
    assert "Mine" not in p and "WhatsApp" in p
    assert client.post("/settings/sources/5/delete").status_code == 404     # not this shop's


def test_past_bills_columns_and_excel_follow_the_shops_methods(client):
    setup(client)
    client.post("/settings/payments", data={"name": "Credit"})
    bill(client, mode="Credit")
    bill(client, mode="Cash")
    past = page(client, "/bills")
    for name in ("Cash", "UPI", "Card", "Bank transfer", "Credit"):
        assert f'<th class="num">{name}</th>' in past
    assert "Net-banking" not in past
    wb = load_workbook(BytesIO(client.get("/bills/export.xlsx").data))
    header = [c.value for c in wb["Bills"][1]]
    for name in ("Cash", "UPI", "Card", "Bank transfer", "Credit", "Other payment", "Due"):
        assert name in header
    row = next(dict(zip(header, [c.value for c in r])) for r in wb["Bills"].iter_rows(min_row=2)
               if r[15 + 4].value)             # the Credit column
    assert row["Credit"] == 450 and row["Total"] == 450


def test_existing_shops_keep_their_old_choices_when_upgraded(app):
    path = app.config["DATABASE"]
    raw = sqlite3.connect(path)
    raw.execute("INSERT INTO shops (name) VALUES ('Old shop')")
    raw.execute("DROP TABLE order_sources")
    raw.execute("DROP TABLE payment_methods")
    raw.commit()
    raw.close()
    with app.app_context():
        init_db()
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    assert choices.names(db, 1, "source") == ["Store (walk-in)", "Stall", "Website"]
    assert choices.names(db, 1, "method") == ["UPI", "Cash", "Net-banking"]


def test_terms_keep_their_line_breaks_on_the_bill(client):
    setup(client)
    client.post("/settings/details", data={"name": "Shop", "terms": "1. No returns\n2. No refunds"})
    bill(client)
    assert "white-space: pre-line" in client.get("/static/css/app.css").data.decode()
    assert "1. No returns<br>2. No refunds" in page(client, "/bills/1")
