import sqlite3
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.db import init_db
from app.services.profit import profit_figures
from tests.test_auth import create


def product(client, name, price, cost="", supplier="", stock="50", category=""):
    for field, word in ((category, "categories"), (supplier, "suppliers")):
        if field:
            client.post(f"/settings/{word}", data={"name": field})
    res = client.post("/inventory/new", data={
        "name": name, "price": price, "stock": stock, "low_stock_level": "2",
        "cost_price": cost, "supplier": supplier, "category": category})
    assert res.status_code == 302


def bill(client, items, **extra):
    body = {"items": items, "customer": "Asha", "order_source": "Instagram",
            "payments": [{"mode": "Cash"}], **extra}
    res = client.post("/billing/save", json=body)
    assert res.status_code == 200, res.get_json()


def figures(app, **kw):
    import sqlite3
    db = sqlite3.connect(app.config["DATABASE"])
    db.row_factory = sqlite3.Row
    return profit_figures(db, 1, **kw)


def row(f, name):
    return next(r for r in f["rows"] if r["name"] == name)


def test_profit_is_sales_minus_cost_per_product(client, app):
    create(client)
    product(client, "Tray", "450", cost="300", supplier="Mehta")
    bill(client, [{"product_id": 1, "quantity": 2}])
    f = figures(app)
    tray = row(f, "Tray")
    assert (tray["revenue"], tray["cost"], tray["profit"]) == (90000, 60000, 30000)
    assert tray["per_item"] == 15000 and tray["qty"] == 2
    assert f["profit"] == 30000 and round(f["margin"], 1) == 33.3


def test_a_bill_discount_is_shared_out_and_shipping_is_left_out(client, app):
    create(client)
    product(client, "Tray", "450", cost="300")
    product(client, "Cushion", "300", cost="100")
    bill(client, [{"product_id": 1, "quantity": 2}, {"product_id": 2, "quantity": 1}],
         whole_discount="10", shipping="80")      # 1200 less 10% = 1080, plus 80 shipping
    f = figures(app)
    assert row(f, "Tray")["revenue"] == 81000 and row(f, "Cushion")["revenue"] == 27000
    assert f["profit"] == 108000 - 70000


def test_items_without_a_cost_are_counted_apart_not_guessed(client, app):
    create(client)
    product(client, "Tray", "450", cost="300")
    product(client, "Mystery", "200")
    bill(client, [{"product_id": 1, "quantity": 1}, {"product_id": 2, "quantity": 1}])
    f = figures(app)
    assert f["profit"] == 15000 and f["unknown_revenue"] == 20000 and f["no_cost_count"] == 1
    page = client.get("/profit").data.decode()
    assert "have no cost price" in page or "has no cost price" in page
    # entering the cost later counts the old sale too
    client.post("/inventory/2/edit", data={"name": "Mystery", "price": "200", "stock": "49",
                                           "low_stock_level": "2", "cost_price": "150"})
    assert figures(app)["profit"] == 15000 + 5000


def test_cost_is_copied_onto_the_bill_when_it_is_made(client, app):
    create(client)
    product(client, "Tray", "450", cost="300")
    bill(client, [{"product_id": 1, "quantity": 1}])
    client.post("/inventory/1/edit", data={"name": "Tray", "price": "450", "stock": "49",
                                           "low_stock_level": "2", "cost_price": "400"})
    assert row(figures(app), "Tray")["profit"] == 15000          # the old bill keeps its cost


def test_cancelled_bills_and_other_dates_do_not_count(client, app):
    create(client)
    product(client, "Tray", "450", cost="300")
    bill(client, [{"product_id": 1, "quantity": 1}])
    bill(client, [{"product_id": 1, "quantity": 1}], order_date=(date.today() - timedelta(days=40)).isoformat())
    client.post("/bills/1/cancel")
    assert figures(app)["profit"] == 15000                       # only the older bill is left
    today = date.today().isoformat()
    assert figures(app, start=today, end=today)["profit"] == 0


def test_unsold_products_are_listed_and_the_page_loads(client, app):
    create(client)
    product(client, "Tray", "450", cost="300", supplier="Mehta")
    product(client, "Cushion", "300")
    page = client.get("/profit").data.decode()
    assert "<h1>Profit</h1>" in page and "Tray" in page and "Cushion" in page
    assert "Mehta" in page and "₹150.00" in page                   # profit per item
    assert ">Profit</a>" in client.get("/home").data.decode()      # menu tab


def test_shops_are_kept_separate(client, app):
    create(client)
    product(client, "Tray", "450", cost="300")
    bill(client, [{"product_id": 1, "quantity": 1}])
    client.get("/logout")
    create(client, login="two@example.com", shop="Other Shop")
    assert "Tray" not in client.get("/profit").data.decode()


# ---- Home: profit and supplier views ----

def test_home_has_profit_and_supplier_views(client):
    create(client)
    product(client, "Tray", "450", cost="300", supplier="Mehta", stock="10")
    product(client, "Cushion", "300", supplier="Rao", stock="10")
    bill(client, [{"product_id": 1, "quantity": 2}])
    page = client.get("/home").data.decode()
    assert "By supplier" in page and "Mehta" in page and "Rao" in page
    assert "₹3,600.00" in page and "₹3,000.00" in page                # stock value by supplier (8 trays left)
    assert "how this month" in page and "₹300.00" in page             # profit view: 900 - 600
    assert "Profit" in page


def test_home_with_no_cost_prices_explains_the_profit_view(client):
    create(client)
    product(client, "Tray", "450")
    page = client.get("/home").data.decode()
    assert "No profit to show yet" in page and "No supplier" in page


# ---- Excel download ----

def sheets(client, url="/bills/export.xlsx"):
    res = client.get(url)
    assert res.status_code == 200
    assert "spreadsheetml" in res.mimetype and "attachment" in res.headers["Content-Disposition"]
    return load_workbook(BytesIO(res.data))


def test_excel_has_a_row_per_bill_and_per_item(client):
    create(client)
    product(client, "Tray", "450", cost="300")
    bill(client, [{"product_id": 1, "quantity": 2}], due_amount="300",
         payments=[{"mode": "Cash", "amount": "400"}, {"mode": "UPI", "amount": "200"}])
    bill(client, [{"product_id": 1, "quantity": 1}], customer="=1+1")
    wb = sheets(client)
    rows = list(wb["Bills"].iter_rows(values_only=True))
    header = rows[0]
    first = next(dict(zip(header, r)) for r in rows[1:] if r[2] == "Asha")
    assert first["Total"] == 900 and first["UPI"] == 200 and first["Cash"] == 400 and first["Due"] == 300
    assert first["Bill no"].endswith("/1")
    items = list(wb["Items"].iter_rows(values_only=True))
    assert len(rows) == 3 and len(items) == 3
    assert {r[2] for r in rows[1:]} == {"Asha", "=1+1"}
    formulas = [c for c in wb["Bills"]["C"] if c.data_type == "f"]
    assert formulas == []                       # a typed '=1+1' is text, never a formula


def test_excel_follows_the_filters_and_needs_sign_in(client):
    assert client.get("/bills/export.xlsx").status_code == 302
    create(client)
    product(client, "Tray", "450", cost="300")
    bill(client, [{"product_id": 1, "quantity": 1}], customer="Asha")
    bill(client, [{"product_id": 1, "quantity": 1}], customer="Ravi")
    wb = sheets(client, "/bills/export.xlsx?q=Ravi")
    assert [r[2] for r in wb["Bills"].iter_rows(min_row=2, values_only=True)] == ["Ravi"]
    page = client.get("/bills").data.decode()
    assert "Download Excel" in page and "/bills/export.xlsx" in page


# ---- the vendor -> supplier rename ----

def test_old_databases_with_a_vendor_column_are_upgraded(app):
    path = app.config["DATABASE"]
    raw = sqlite3.connect(path)
    raw.execute("ALTER TABLE products RENAME COLUMN supplier TO vendor")
    raw.commit()
    raw.close()
    with app.app_context():
        init_db()
    columns = [r[1] for r in sqlite3.connect(path).execute("PRAGMA table_info(products)")]
    assert "supplier" in columns and "vendor" not in columns


def test_profit_excel_matches_the_page_and_follows_the_dates(client):
    create(client)
    product(client, "Tray", "450", cost="300", supplier="Mehta")
    product(client, "Mystery", "200")
    bill(client, [{"product_id": 1, "quantity": 2}, {"product_id": 2, "quantity": 1}])
    assert "/profit/export.xlsx" in client.get("/profit").data.decode()
    res = client.get("/profit/export.xlsx")
    assert res.status_code == 200 and "spreadsheetml" in res.mimetype
    rows = list(load_workbook(BytesIO(res.data))["Profit"].iter_rows(values_only=True))
    header = rows[2]
    tray = dict(zip(header, next(r for r in rows if r[0] == "Tray")))
    assert (tray["Supplier"], tray["Sales"], tray["Cost"], tray["Profit"]) == ("Mehta", 900, 600, 300)
    total = dict(zip(header, next(r for r in rows if r[0] == "Total")))
    assert total["Profit"] == 300 and total["Quantity sold"] == 3
    assert any(r[0] and "not counted" in str(r[0]) for r in rows)
    old = client.get("/profit/export.xlsx?date_from=2000-01-01&date_to=2000-01-02")
    rows = list(load_workbook(BytesIO(old.data))["Profit"].iter_rows(values_only=True))
    assert dict(zip(header, next(r for r in rows if r[0] == "Total")))["Profit"] == 0
    client.get("/logout")
    assert client.get("/profit/export.xlsx").status_code == 302
