"""Shop Details email, defaults, categories and suppliers; Customer Details filters; calendar
defaults; Inventory labels and filters."""
import html
import io
import re
import sqlite3
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.db import init_db
from tests.test_auth import create
from tests.test_inventory import add


def page(client, url):
    return html.unescape(client.get(url).data.decode())


def page_of(res):
    return html.unescape(res.data.decode())


def post_bill(client, name="Asha", phone="", **extra):
    body = {"items": [{"product_id": 1, "quantity": 1}], "customer": name, "phone": phone,
            "order_source": "Walk-in", "payments": [{"mode": "Cash"}], **extra}
    res = client.post("/billing/save", json=body)
    assert res.status_code == 200, res.get_json()


def setup(client):
    create(client)
    add(client, name="Tray", price="450", stock="500")


def ago(days):
    return (date.today() - timedelta(days=days)).isoformat()


def first_of_month():
    return date.today().replace(day=1).isoformat()


# ---- Shop Details ----

def test_email_is_saved_shown_in_the_preview_and_printed_on_the_bill(client):
    setup(client)
    res = client.post("/settings/details", data={"name": "Shop", "email": "hello@shop.example"},
                      follow_redirects=True)
    p = page_of(res)
    assert 'value="hello@shop.example"' in p and 'id="pv-email">hello@shop.example<' in p
    post_bill(client)
    assert "hello@shop.example" in page(client, "/bills/1")
    bad = client.post("/settings/details", data={"name": "Shop", "email": "not an email"},
                      follow_redirects=True)
    assert "doesn't look right" in page_of(bad) and 'value="hello@shop.example"' in page(client, "/settings")


def test_a_default_order_source_and_payment_method_start_new_bills(client):
    setup(client)
    assert "<option selected>" not in page(client, "/billing").split('id="source"')[1].split("</select>")[0]
    client.post("/settings/sources/default", data={"name": "WhatsApp"})
    client.post("/settings/payments/default", data={"name": "UPI"})
    billing = page(client, "/billing")
    assert "<option selected>WhatsApp</option>" in billing and "<option selected>UPI</option>" in billing
    settings = page(client, "/settings")
    assert "Default value" in settings and "<option selected>WhatsApp</option>" in settings
    post_bill(client, order_source="WhatsApp", payments=[{"mode": "UPI"}])
    # editing a saved bill shows that bill's own choices, not the defaults
    edit = page(client, "/billing/edit/1")
    assert "<option selected>WhatsApp</option>" not in edit.split('id="source"')[1].split("</select>")[0]


def test_default_must_be_in_the_list_and_goes_away_when_removed(client):
    setup(client)
    client.post("/settings/sources/default", data={"name": "WhatsApp"})
    res = client.post("/settings/sources/default", data={"name": "Carrier pigeon"}, follow_redirects=True)
    assert "choose one of the order sources" in page_of(res)
    assert "<option selected>WhatsApp</option>" in page(client, "/billing")
    ids = {n: i for n, i in re.findall(r'<span>([^<]+)</span>\s*<form method="post" action="/settings/sources/(\d+)/delete"', page(client, "/settings"))}
    client.post(f"/settings/sources/{ids['WhatsApp']}/delete")
    assert "<option selected>" not in page(client, "/billing").split('id="source"')[1].split("</select>")[0]
    client.post("/settings/sources/default", data={"name": "Walk-in"})
    client.post("/settings/sources/default", data={"name": ""})                 # back to none
    assert "<option selected>" not in page(client, "/billing").split('id="source"')[1].split("</select>")[0]


def test_categories_and_suppliers_are_lists_in_shop_details(client):
    create(client)
    p = page(client, "/settings")
    for text in ("Product categories", "Suppliers", "Add a category", "Add a supplier"):
        assert text in p
    client.post("/settings/categories", data={"name": "Bedsheets"})
    client.post("/settings/suppliers", data={"name": "Mehta Traders"})
    p = page(client, "/settings")
    assert "Bedsheets" in p and "Mehta Traders" in p
    dup = client.post("/settings/categories", data={"name": "bedsheets"}, follow_redirects=True)
    assert "already in the list" in page_of(dup)
    # "Other" is only special for bills, so it is a fine category
    assert "added" in page_of(client.post("/settings/categories", data={"name": "Other"}, follow_redirects=True))
    ids = re.findall(r'/settings/categories/(\d+)/delete', page(client, "/settings"))
    assert client.post(f"/settings/categories/{ids[0]}/delete").status_code == 302


def test_lists_stay_private_to_each_shop(client):
    create(client)
    client.post("/settings/categories", data={"name": "Lampshades"})
    client.get("/logout")
    create(client, login="two@example.com", shop="Other")
    assert "Lampshades" not in page(client, "/settings")
    assert client.post("/settings/categories/1/delete").status_code == 404


def test_the_product_form_uses_dropdowns_from_those_lists(client):
    create(client)
    p = page(client, "/inventory/new")
    assert "Add categories in" in p and "Add suppliers in" in p
    client.post("/settings/categories", data={"name": "Bedsheets"})
    client.post("/settings/suppliers", data={"name": "Mehta Traders"})
    p = page(client, "/inventory/new")
    assert '<select name="category">' in p and "<option>Bedsheets</option>" in p
    assert '<select name="supplier">' in p and "<option>Mehta Traders</option>" in p
    ok = client.post("/inventory/new", data={"name": "Sheet", "price": "100", "stock": "5",
                                             "low_stock_level": "1", "category": "Bedsheets",
                                             "supplier": "Mehta Traders"})
    assert ok.status_code == 302
    assert "<option selected>Bedsheets</option>" in page(client, "/inventory/1/edit")
    bad = client.post("/inventory/new", data={"name": "X", "price": "1", "stock": "1", "low_stock_level": "1",
                                              "category": "Made up"})
    assert bad.status_code == 200 and "choose a category from the list" in page_of(bad)
    bad = client.post("/inventory/new", data={"name": "X", "price": "1", "stock": "1", "low_stock_level": "1",
                                              "supplier": "Nobody"})
    assert "choose a supplier from the list" in page_of(bad)
    none = client.post("/inventory/new", data={"name": "Plain", "price": "1", "stock": "1", "low_stock_level": "1"})
    assert none.status_code == 302                                        # no category is fine


def test_a_product_can_keep_a_category_the_list_no_longer_has(client):
    create(client)
    client.post("/settings/categories", data={"name": "Bedsheets"})
    client.post("/inventory/new", data={"name": "Sheet", "price": "100", "stock": "5",
                                        "low_stock_level": "1", "category": "Bedsheets"})
    cid = re.findall(r'/settings/categories/(\d+)/delete', page(client, "/settings"))[0]
    client.post(f"/settings/categories/{cid}/delete")
    edit = page(client, "/inventory/1/edit")
    assert "<option selected>Bedsheets</option>" in edit                  # still shown, still kept
    res = client.post("/inventory/1/edit", data={"name": "Sheet", "price": "120", "stock": "5",
                                                 "low_stock_level": "1", "category": "Bedsheets"})
    assert res.status_code == 302


def test_existing_products_give_their_categories_and_suppliers_to_the_new_lists(app):
    path = app.config["DATABASE"]
    raw = sqlite3.connect(path)
    raw.execute("INSERT INTO shops (name) VALUES ('Old')")
    raw.execute("INSERT INTO products (shop_id, name, category, supplier, price) VALUES (1, 'A', 'Trays', 'Mehta', 100)")
    raw.execute("INSERT INTO products (shop_id, name, category, supplier, price) VALUES (1, 'B', 'Jars', 'Mehta', 100)")
    raw.execute("DROP TABLE product_categories")
    raw.execute("DROP TABLE suppliers")
    raw.commit()
    raw.close()
    with app.app_context():
        init_db()
    db = sqlite3.connect(path)
    assert [r[0] for r in db.execute("SELECT name FROM product_categories ORDER BY position")] == ["Jars", "Trays"]
    assert [r[0] for r in db.execute("SELECT name FROM suppliers")] == ["Mehta"]


# ---- Inventory ----

def test_inventory_labels_supplier_column_filters_and_inactive_button(client):
    create(client)
    add(client, name="Tray", category="Trays", supplier="Mehta", price="450", code="KA36")
    add(client, name="Sheet", category="Bedsheets", supplier="Rao", price="300")
    p = page(client, "/inventory/")
    for text in ("Item Code/Bar Code", "Selling Price", "Supplier", "Mehta", "Rao"):
        assert text in p
    assert ">Inactive<" in p and ">Remove<" not in p
    form = page(client, "/inventory/new")
    assert "Item Code/Bar Code (optional)" in form and "Selling Price (₹)" in form and "Price (₹)" not in form.replace("Selling Price (₹)", "")
    # filters
    assert "Trays</option>" in p and "Rao</option>" in p                        # the filter dropdowns
    by_category = page(client, "/inventory/?category=Trays")
    assert "<td>Tray</td>" in by_category and "<td>Sheet</td>" not in by_category
    by_supplier = page(client, "/inventory/?supplier=Rao")
    assert "<td>Sheet</td>" in by_supplier and "<td>Tray</td>" not in by_supplier
    both = page(client, "/inventory/?category=Trays&supplier=Rao")
    assert "<td>Tray</td>" not in both and "<td>Sheet</td>" not in both
    assert "<td>Tray</td>" in page(client, "/inventory/?category=Nonsense")           # unknown = no filter


def test_inactive_hides_the_product_but_its_code_stays_used(client):
    create(client)
    add(client, name="Tray", code="KA36")
    client.post("/inventory/1/remove")
    assert "<td>Tray</td>" not in page(client, "/inventory/")
    again = add(client, name="Other", code="KA36")
    assert again.status_code == 200 and "has been removed" in page_of(again)


# ---- Customer Details ----

def test_customers_default_to_last_bill_first_and_can_be_sorted_by_name(client, app):
    setup(client)
    post_bill(client, "Anna", "9000000001")
    post_bill(client, "Zed", "9000000002")
    db = sqlite3.connect(app.config["DATABASE"])
    today = date.today().isoformat()
    db.execute("UPDATE bills SET created_at = ? WHERE id = 1", (today + "T01:00:00",))   # Anna: earlier
    db.execute("UPDATE bills SET created_at = ? WHERE id = 2", (today + "T23:00:00",))   # Zed: later
    db.commit()
    names = lambda url: re.findall(r'<td><a href="/customers/\d+">([^<]+)</a></td>', page(client, url))
    assert names("/customers") == ["Zed", "Anna"]                           # newest bill first
    assert names("/customers?sort=name") == ["Anna", "Zed"]
    assert 'value="last" selected' in page(client, "/customers")


def test_customers_can_be_filtered_by_retail_or_b2b(client):
    setup(client)
    post_bill(client, "Shop One", "9000000001", customer_type="B2B")
    post_bill(client, "Person", "9000000002", customer_type="Retail")
    names = lambda url: re.findall(r'<td><a href="/customers/\d+">([^<]+)</a></td>', page(client, url))
    assert sorted(names("/customers")) == ["Person", "Shop One"]
    assert names("/customers?type=B2B") == ["Shop One"]
    assert names("/customers?type=Retail") == ["Person"]
    assert sorted(names("/customers?type=Bogus")) == ["Person", "Shop One"]


def test_customers_have_a_calendar_filter_and_start_with_this_month(client):
    setup(client)
    post_bill(client, "Old Customer", "9000000001", order_date=ago(60))
    post_bill(client, "New Customer", "9000000002")
    names = lambda url: re.findall(r'<td><a href="/customers/\d+">([^<]+)</a></td>', page(client, url))
    assert names("/customers") == ["New Customer"]                          # this month so far
    p = page(client, "/customers")
    assert f'name="date_from" value="{first_of_month()}"' in p and f'name="date_to" value="{date.today().isoformat()}"' in p
    assert sorted(names("/customers?date_from=&date_to=")) == ["New Customer", "Old Customer"]
    assert names(f"/customers?date_from={ago(90)}&date_to={ago(30)}") == ["Old Customer"]
    # counts and totals are for the period only
    post_bill(client, "Old Customer", "9000000001")
    row = page(client, f"/customers?date_from={ago(90)}&date_to={ago(30)}")
    assert ">1</a>" in row                                                  # one bill in that period
    assert ">2</a>" in page(client, "/customers?date_from=&date_to=")


# ---- calendars everywhere start at the 1st of the month ----

def test_past_bills_start_with_this_month_and_the_excel_follows(client):
    setup(client)
    post_bill(client, "Old Bill Customer", order_date=ago(60))
    post_bill(client, "New Bill Customer")
    p = page(client, "/bills")
    assert "New Bill Customer" in p and "Old Bill Customer" not in p
    assert f'name="date_from" value="{first_of_month()}"' in p
    both = page(client, "/bills?date_from=&date_to=")
    assert "New Bill Customer" in both and "Old Bill Customer" in both
    names = lambda res: [r[2] for r in load_workbook(BytesIO(res.data))["Bills"].iter_rows(min_row=2, values_only=True)]
    assert names(client.get("/bills/export.xlsx")) == ["New Bill Customer"]
    assert sorted(names(client.get("/bills/export.xlsx?date_from=&date_to="))) == ["New Bill Customer", "Old Bill Customer"]
    assert 'href="/bills"' in p                                              # a way back after filtering
    assert "/bills/export.xlsx" in p


def test_profit_starts_with_this_month_and_its_excel_follows(client):
    setup(client)
    client.post("/settings/categories", data={"name": "x"})
    client.post("/inventory/1/edit", data={"name": "Tray", "price": "450", "stock": "500",
                                           "low_stock_level": "3", "cost_price": "300"})
    post_bill(client, "Old", order_date=ago(60))
    post_bill(client, "New")
    p = page(client, "/profit")
    assert f'name="date_from" value="{first_of_month()}"' in p
    assert "₹150.00" in p                                                    # one bill's profit, not two
    wide = page(client, "/profit?date_from=&date_to=")
    assert "₹300.00" in wide
    sheet = load_workbook(BytesIO(client.get("/profit/export.xlsx").data))["Profit"]
    assert any(r[0] == "Total" and r[9] == 150 for r in sheet.iter_rows(values_only=True))
    assert any(r[0] == "Total" and r[9] == 300 for r in load_workbook(BytesIO(
        client.get("/profit/export.xlsx?date_from=&date_to=").data))["Profit"].iter_rows(values_only=True))


def test_the_shop_details_page_has_no_unfilled_template_text(client):
    create(client)
    p = client.get("/settings").data.decode()
    assert "{{" not in p and "}}" not in p
    assert "Up to 3 QR codes print on the bill" in html.unescape(p)
